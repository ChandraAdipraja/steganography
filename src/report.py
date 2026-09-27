"""Builder laporan XLSX bersama — dipakai scripts/benchmark.py (CLI)
dan route web /analysis/dataset. Satu sumber kebenaran untuk format tabel.

Kontrak per_cover (list of dict):
    cover, resolution, hist_file, hist_path, r_diff, g_diff, b_diff,
    verdict, lsb_file, jpeg_file, jpeg_success, jpeg_error
Kontrak rows (list of dict):
    image, resolution, message_size, payload_size, mse, psnr, extraction_ok
"""

from __future__ import annotations

import datetime

THUMB_WIDTH = 360
THUMB_HEIGHT = 120


def fmt_psnr(value) -> str:
    if value == float("inf"):
        return "inf"
    return f"{value:.2f}"


def channel_maxdiff(hist) -> int:
    cover_counts = hist["cover"]
    stego_counts = hist["stego"]
    return max(abs(c - s) for c, s in zip(cover_counts, stego_counts))


def build_xlsx(
    rows,
    per_cover,
    dest,
    *,
    cover_desc: str,
    message_desc: str,
    password_desc: str,
    jpeg_quality: int,
    capacity_note: str = "budget plaintext (Opsi A, FULL_OVERHEAD=52)",
):
    """Tulis workbook ke path file atau file-like object (mis. BytesIO)."""
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl belum terinstal. Jalankan: python -m pip install openpyxl"
        ) from exc

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="334155")

    def style_header(ws, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

    # ---- Sheet 1: Ringkasan ----
    ws = wb.active
    ws.title = "Ringkasan"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70
    summary = [
        ("Tanggal run", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Citra cover", cover_desc),
        ("Ukuran pesan", message_desc),
        ("Password", password_desc),
        ("Overhead", "8 B header STG1 + 44 B crypto (salt+nonce+tag)"),
        ("Kapasitas", capacity_note),
        ("JPEG quality", str(jpeg_quality)),
        ("Total sampel PSNR/MSE", str(len(rows))),
    ]
    ws.append(["Parameter", "Nilai"])
    for k, v in summary:
        ws.append([k, v])
    style_header(ws, 2)

    # ---- Sheet 2: PSNR/MSE ----
    ws2 = wb.create_sheet("PSNR_MSE")
    headers2 = ["Citra", "Resolusi", "Pesan (B)", "Blob+overhead (B)",
                "MSE", "PSNR (dB)", "Extraction OK"]
    ws2.append(headers2)
    for r in rows:
        ws2.append([r["image"], r["resolution"], r["message_size"],
                    r["payload_size"], round(r["mse"], 6),
                    fmt_psnr(r["psnr"]), r["extraction_ok"]])
    sizes = sorted({r["message_size"] for r in rows})
    if sizes:
        base = len(rows) + 2
        ws2.cell(row=base, column=1,
                 value="Rata-rata per ukuran pesan").font = Font(bold=True)
        for size in sizes:
            subset = [r for r in rows if r["message_size"] == size]
            avg_mse = sum(r["mse"] for r in subset) / len(subset)
            vals = [r["psnr"] for r in subset if r["psnr"] != float("inf")]
            avg_psnr = sum(vals) / len(vals) if vals else float("inf")
            ws2.append([f"avg pesan {size}B", "", size, "",
                        round(avg_mse, 6), fmt_psnr(avg_psnr), ""])
    for col, w in zip("ABCDEFG", [26, 12, 11, 17, 12, 12, 15]):
        ws2.column_dimensions[col].width = w
    style_header(ws2, 7)

    # ---- Sheet 3: Histogram + thumbnail ----
    ws3 = wb.create_sheet("Histogram")
    headers3 = ["Citra", "Resolusi", "R maxDbin", "G maxDbin", "B maxDbin",
                "Verdict", "File", "Thumbnail"]
    ws3.append(headers3)
    for i, pc in enumerate(per_cover, start=2):
        ws3.append([pc["cover"], pc["resolution"], pc["r_diff"],
                    pc["g_diff"], pc["b_diff"], pc["verdict"],
                    pc["hist_file"], ""])
        # embed file asli, tampil kecil (openpyxl baca file saat save)
        xl_img = XLImage(pc["hist_path"])
        xl_img.width, xl_img.height = THUMB_WIDTH, THUMB_HEIGHT
        ws3.add_image(xl_img, f"H{i}")
        ws3.row_dimensions[i].height = THUMB_HEIGHT * 0.78
    for col, w in zip("ABCDEFGH", [26, 12, 11, 11, 11, 24, 30, 52]):
        ws3.column_dimensions[col].width = w
    style_header(ws3, 8)

    # ---- Sheet 4: JPEG ----
    ws4 = wb.create_sheet("JPEG_Q90")
    headers4 = ["Citra", "Quality", "Success", "Error", "File"]
    ws4.append(headers4)
    for pc in per_cover:
        ws4.append([pc["cover"], jpeg_quality, pc["jpeg_success"],
                    pc["jpeg_error"], pc["jpeg_file"]])
    for col, w in zip("ABCDE", [26, 10, 10, 45, 24]):
        ws4.column_dimensions[col].width = w
    style_header(ws4, 5)

    # ---- Sheet 5: Steganalisis LSB (cover vs stego + bidang LSB) ----
    # Kontrak tambahan per pc: cover_path, stego_file, stego_path,
    #   lsb_cover_file, lsb_cover_path, lsb_stego_file, lsb_stego_path, mse, psnr
    ws5 = wb.create_sheet("Steganalisis_LSB")
    headers5 = ["Citra", "Cover", "Stego", "LSB Cover", "LSB Stego",
                "Observasi", "File"]
    ws5.append(headers5)
    IMG = 140
    for i, pc in enumerate(per_cover, start=2):
        mse_v, psnr_v = pc.get("mse"), pc.get("psnr")
        if mse_v is None or psnr_v is None:
            obs = f"medium dilewati — {pc.get('verdict', '-')}"
        else:
            obs = (f"MSE {mse_v:.6f}, PSNR {fmt_psnr(psnr_v)} dB — "
                   f"{pc.get('verdict', '-')}")
        ws5.append([pc["cover"], "", "", "", "", obs,
                    f"{pc.get('cover_file', pc['cover'])} / "
                    f"{pc.get('stego_file', '-')}"])
        for col_letter, key in (("B", "cover_path"), ("C", "stego_path"),
                                ("D", "lsb_cover_path"), ("E", "lsb_stego_path")):
            path = pc.get(key)
            if path:
                try:
                    thumb = XLImage(path)
                    thumb.width, thumb.height = IMG, IMG
                    ws5.add_image(thumb, f"{col_letter}{i}")
                except Exception:
                    pass
        ws5.row_dimensions[i].height = IMG * 0.78
    for col, w in zip("ABCDEFG", [26, 22, 22, 22, 22, 34, 40]):
        ws5.column_dimensions[col].width = w
    style_header(ws5, 7)

    wb.save(dest)
